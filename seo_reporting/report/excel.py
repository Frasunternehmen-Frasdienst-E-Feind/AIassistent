"""Excel-Export des Reports (.xlsx) – kompakte Arbeitsmappe für das Marketing-Team.

Spiegelt die Werte des JSON-Reports (:mod:`seo_reporting.report.build`) in mehrere
Tabellenblätter: Übersicht, GSC, GA4, CRM und Cluster-Bewegungen. Es werden nur die
bereits berechneten Werte geschrieben (Snapshot), keine Formeln.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

_HEADER_FILL = PatternFill("solid", fgColor="3A4145")
_HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF")
_TITLE_FONT = Font(name="Arial", bold=True, size=13)
_BODY_FONT = Font(name="Arial")


def _autosize(ws: Worksheet, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = width


def _table(ws: Worksheet, top_row: int, headers: list[str], rows: list[list]) -> int:
    """Schreibt eine Tabelle ab ``top_row``; gibt die nächste freie Zeile zurück."""
    for c, head in enumerate(headers, start=1):
        cell = ws.cell(row=top_row, column=c, value=head)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r, values in enumerate(rows, start=top_row + 1):
        for c, value in enumerate(values, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = _BODY_FONT
    return top_row + 1 + len(rows) + 1  # eine Leerzeile Abstand


def _title(ws: Worksheet, text: str, row: int = 1) -> int:
    cell = ws.cell(row=row, column=1, value=text)
    cell.font = _TITLE_FONT
    return row + 2


def _delta_rows(cmp: dict | None, metrics: list[tuple[str, str]], yoy: dict | None) -> list[list]:
    """Baut Zeilen [Kennzahl, Aktuell, Vorperiode, Δ abs, Δ %, YoY Δ %] aus Delta-Dicts."""
    rows: list[list] = []
    for key, label in metrics:
        d = (cmp or {}).get(key)
        if d is None:
            rows.append([label, None, None, None, None, None])
            continue
        yoy_pct = (yoy or {}).get(key, {}).get("pct") if yoy else None
        rows.append([label, d["current"], d["previous"], d["abs"], d["pct"], yoy_pct])
    return rows


def write_xlsx(report: dict, path: Path) -> Path:
    """Schreibt ``report`` als .xlsx nach ``path`` und gibt den Pfad zurück."""
    meta = report["meta"]
    periods = meta["periods"]
    prev_label = periods["previous_label"]
    wb = Workbook()

    # ---- Übersicht --------------------------------------------------------
    ws = wb.active
    ws.title = "Übersicht"
    ws.sheet_view.showGridLines = False
    row = _title(ws, "SEO-/Lead-Report – Übersicht")
    info = [
        ["Tool", meta["tool"]],
        ["Workflow", meta["workflow"]],
        ["Website", meta["site"]],
        ["Zeitraum", f"{periods['current']['label']} ({periods['current']['start']} bis {periods['current']['end']})"],
        ["Vergleich", f"{prev_label}: {periods['previous']['label']}"],
        ["YoY", periods["yoy"]["label"] if report["gsc"]["totals"]["YoY"] is not None else "–"],
        ["Erstellt", meta["generated_at"]],
    ]
    row = _table(ws, row, ["Feld", "Wert"], info)
    notes = meta.get("data_notes") or []
    row = _table(ws, row, ["Datenhinweise"], [[n] for n in notes] or [["keine"]])
    _autosize(ws, {1: 22, 2: 60})

    # ---- GSC --------------------------------------------------------------
    ws = wb.create_sheet("GSC")
    ws.sheet_view.showGridLines = False
    row = _title(ws, "Google Search Console")
    totals = report["gsc"]["totals"]
    metrics = [("clicks", "Klicks"), ("impressions", "Impressionen"), ("ctr", "CTR (%)"), ("position", "Ø Position")]
    row = _table(
        ws, row,
        ["Kennzahl", "Aktuell", f"Vorperiode ({prev_label})", "Δ abs", f"Δ % ({prev_label})", "Δ % (YoY)"],
        _delta_rows(totals.get(prev_label), metrics, totals.get("YoY")),
    )
    clusters = report["gsc"]["clusters"]["current"]
    row = _table(
        ws, row,
        ["Cluster", "Klicks", "Impressionen", "CTR", "Ø Position"],
        [[c["name"], c["clicks"], c["impressions"], c["ctr"], c["position"]] for c in clusters],
    )
    row = _table(
        ws, row,
        ["Top-Query", "Klicks", "Impressionen", "CTR", "Ø Position"],
        [[q["key"], q["clicks"], q["impressions"], q["ctr"], q["position"]] for q in report["gsc"]["top_queries"]],
    )
    row = _table(
        ws, row,
        ["Top-Seite", "Klicks", "Impressionen", "CTR", "Ø Position"],
        [[p["key"], p["clicks"], p["impressions"], p["ctr"], p["position"]] for p in report["gsc"]["top_pages"]],
    )
    _autosize(ws, {1: 46, 2: 12, 3: 16, 4: 10, 5: 12, 6: 12})

    # ---- GA4 --------------------------------------------------------------
    ws = wb.create_sheet("GA4")
    ws.sheet_view.showGridLines = False
    row = _title(ws, "Google Analytics 4")
    g_totals = report["ga4"]["totals"]
    g_metrics = [("sessions", "Sitzungen"), ("engaged_sessions", "Engagierte Sitzungen"),
                 ("total_users", "Nutzer"), ("key_events", "Schlüsselereignisse")]
    row = _table(
        ws, row,
        ["Kennzahl", "Aktuell", f"Vorperiode ({prev_label})", "Δ abs", f"Δ % ({prev_label})", "Δ % (YoY)"],
        _delta_rows(g_totals.get(prev_label), g_metrics, g_totals.get("YoY")),
    )
    ch_rows = []
    for c in report["ga4"]["channels"]:
        cur = c["current"]
        d = c.get(prev_label) or {}
        ch_rows.append([c["channel"], cur["sessions"], cur["engaged_sessions"], cur["total_users"],
                        cur["key_events"], d.get("sessions", {}).get("pct")])
    row = _table(ws, row, ["Kanal", "Sitzungen", "Engagiert", "Nutzer", "Schlüsselereignisse", f"Δ % Sitzungen ({prev_label})"], ch_rows)
    row = _table(
        ws, row,
        ["Landingpage", "Sitzungen", "Engagiert", "Nutzer", "Schlüsselereignisse"],
        [[lp["key"], lp["sessions"], lp["engaged_sessions"], lp["total_users"], lp["key_events"]]
         for lp in report["ga4"]["landing_pages"]],
    )
    _autosize(ws, {1: 40, 2: 12, 3: 12, 4: 10, 5: 18, 6: 22})

    # ---- CRM --------------------------------------------------------------
    ws = wb.create_sheet("CRM")
    ws.sheet_view.showGridLines = False
    row = _title(ws, "CRM / Leads")
    crm = report["crm"]["current"]
    if not crm.get("available"):
        row = _table(ws, row, ["Hinweis"], [[crm.get("note") or "Keine CRM-Daten verfügbar."]])
    else:
        cmp = report["crm"].get(prev_label) or {}
        summary = [
            ["Leads", crm["leads"], (cmp.get("leads") or {}).get("pct")],
            ["Gewonnen", crm["won"], (cmp.get("won") or {}).get("pct")],
            ["Wert (EUR)", crm["value"], (cmp.get("value") or {}).get("pct")],
        ]
        row = _table(ws, row, ["Kennzahl", "Aktuell", f"Δ % ({prev_label})"], summary)
        row = _table(
            ws, row,
            ["Quelle", "Leads", "Gewonnen", "Wert (EUR)"],
            [[s["source"], s["leads"], s["won"], s["value"]] for s in crm.get("by_source", [])],
        )
    _autosize(ws, {1: 22, 2: 14, 3: 16, 4: 14})

    # ---- Bewegungen -------------------------------------------------------
    ws = wb.create_sheet("Bewegungen")
    ws.sheet_view.showGridLines = False
    row = _title(ws, "Cluster-Bewegungen (Auto-Highlights)")
    mv_rows = []
    for m in report["movements"]:
        change = m.get("delta_pct")
        change = f"{change:+g} %" if change is not None else (f"{m['delta_abs']:+g}" if m.get("delta_abs") is not None else "–")
        mv_rows.append([m.get("direction"), m.get("cluster"), m.get("basis"), m.get("severity"),
                        m.get("metric"), m.get("previous"), m.get("current"), change])
    row = _table(ws, row, ["Richtung", "Cluster", "Basis", "Schwere", "Metrik", "Vorher", "Nachher", "Δ"],
                 mv_rows or [["–", "keine Bewegung über Schwellenwerten", "", "", "", None, None, ""]])
    _autosize(ws, {1: 10, 2: 26, 3: 12, 4: 10, 5: 10, 6: 12, 7: 12, 8: 14})

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
