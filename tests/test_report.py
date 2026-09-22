import json
from datetime import date
from pathlib import Path

from seo_reporting.cli import main
from seo_reporting.config import load_config
from seo_reporting.periods import periods_for
from seo_reporting.report.build import build_report
from seo_reporting.report.markdown import render
from seo_reporting.sources.demo import DemoSources

ROOT = Path(__file__).resolve().parents[1]


def test_build_and_render_monthly_demo():
    cfg = load_config(ROOT / "config.example.yaml")
    report = build_report(cfg, periods_for("monthly", date(2026, 9, 1)), DemoSources(cfg))
    json.dumps(report)  # muss serialisierbar sein
    assert report["meta"]["periods"]["current"]["label"] == "2026-08"
    assert report["gsc"]["totals"]["YoY"] is not None
    assert report["crm"]["current"]["available"]
    names = [c["name"] for c in report["gsc"]["clusters"]["current"]]
    assert names[0].startswith("Marke") and names[-1].startswith("Sonstige")
    md = render(report)
    assert "# Monatsreport (Workflow 1)" in md
    assert "## 4. CRM / Leads" in md
    assert "Datenhinweise" in md  # Demo simuliert GSC-Verzögerung


def test_weekly_has_no_yoy_by_default():
    cfg = load_config(ROOT / "config.example.yaml")
    report = build_report(cfg, periods_for("weekly", date(2026, 9, 7)), DemoSources(cfg))
    assert report["gsc"]["totals"]["YoY"] is None
    assert "WoW" in report["gsc"]["totals"]
    assert "YoY" not in render(report).split("## 2.")[0]


def test_cli_writes_files(tmp_path, monkeypatch):
    monkeypatch.setenv("REPORT_OUTPUT_DIR", str(tmp_path))
    rc = main(["--config", str(ROOT / "config.example.yaml"), "monthly", "--demo", "--as-of", "2026-09-01"])
    assert rc == 0
    files = sorted(p.name for p in tmp_path.iterdir())
    assert files == ["monthly_2026-08-01_2026-08-31_demo.json", "monthly_2026-08-01_2026-08-31_demo.md"]
    data = json.loads((tmp_path / files[0]).read_text(encoding="utf-8"))
    assert data["meta"]["workflow"] == "monthly"


def test_config_resolution_env_and_missing(tmp_path, monkeypatch):
    from seo_reporting.cli import _resolve_config_path

    monkeypatch.chdir(tmp_path)
    monkeypatch.delenv("SEO_REPORTING_CONFIG", raising=False)
    assert _resolve_config_path(None) is None  # nichts im Ordner → Standardwerte
    monkeypatch.setenv("SEO_REPORTING_CONFIG", str(ROOT / "config.example.yaml"))
    assert _resolve_config_path(None) == str(ROOT / "config.example.yaml")
    assert _resolve_config_path("explizit.yaml") == "explizit.yaml"  # --config gewinnt


def test_missing_config_file_raises():
    import pytest

    with pytest.raises(FileNotFoundError):
        load_config("/gibt/es/nicht.yaml")


def test_console_script_entry_point_defined():
    import tomllib

    data = tomllib.loads((ROOT / "pyproject.toml").read_text(encoding="utf-8"))
    assert data["project"]["scripts"]["seo-report"] == "seo_reporting.cli:main"


def test_write_xlsx_demo(tmp_path):
    from openpyxl import load_workbook

    from seo_reporting.report.excel import write_xlsx

    cfg = load_config(ROOT / "config.example.yaml")
    report = build_report(cfg, periods_for("monthly", date(2026, 9, 1)), DemoSources(cfg))
    out = write_xlsx(report, tmp_path / "report.xlsx")
    assert out.exists()
    wb = load_workbook(out)
    assert wb.sheetnames == ["Übersicht", "GSC", "GA4", "CRM", "Bewegungen"]
    assert wb["Übersicht"]["A1"].value == "SEO-/Lead-Report – Übersicht"
    # Website-Zeile spiegelt die Meta-Angabe
    first_col = [row[0] for row in wb["Übersicht"].iter_rows(values_only=True)]
    assert "Website" in first_col


def test_cli_xlsx_flag(tmp_path, monkeypatch):
    monkeypatch.setenv("REPORT_OUTPUT_DIR", str(tmp_path))
    rc = main(["--config", str(ROOT / "config.example.yaml"), "weekly", "--demo", "--as-of", "2026-09-07", "--xlsx"])
    assert rc == 0
    names = sorted(p.name for p in tmp_path.iterdir())
    assert any(n.endswith(".xlsx") for n in names)
    assert any(n.endswith(".md") for n in names) and any(n.endswith(".json") for n in names)


def test_write_xlsx_escapes_formula_like_values(tmp_path):
    from openpyxl import load_workbook

    from seo_reporting.report.excel import write_xlsx

    cfg = load_config(ROOT / "config.example.yaml")
    report = build_report(cfg, periods_for("monthly", date(2026, 9, 1)), DemoSources(cfg))
    # Formelartige (potenziell bösartige) Suchanfrage einschleusen
    report["gsc"]["top_queries"][0]["key"] = "=1+1"
    out = write_xlsx(report, tmp_path / "report.xlsx")
    values = [v for row in load_workbook(out)["GSC"].iter_rows(values_only=True) for v in row]
    assert "'=1+1" in values            # neutralisiert, als Text gespeichert
    assert "=1+1" not in values          # nie als aktive Formel


def test_email_mime_for():
    from seo_reporting.report.email import _mime_for

    assert _mime_for(Path("r.xlsx")) == ("application", "vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert _mime_for(Path("r.json")) == ("application", "json")
    assert _mime_for(Path("r.md")) == ("text", "markdown")
